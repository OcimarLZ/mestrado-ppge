import configparser
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import create_engine, text


UF_SIGLAS = {
    "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT",
    "MS", "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO",
    "RR", "SC", "SP", "SE", "TO",
}


def carregar_engine() -> object:
    root = Path(__file__).resolve().parents[2]
    config_path = root / "config.ini"
    default_sqlite = root / "INEP.db"

    if not config_path.exists():
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    parser = configparser.ConfigParser()
    parser.read(config_path, encoding="utf-8")

    if "BD_INEP" not in parser:
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    bd = parser["BD_INEP"]
    sgdb = bd.get("SGDB", "sqlite").strip().lower()

    if sgdb == "sqlite":
        pasta_local = bd.get("PASTA_LOCAL", "").replace("\\", "/")
        arquivo = bd.get("DW", "INEP.db").strip()

        if pasta_local:
            pasta_limpa = pasta_local.strip("/")
            sqlite_path = Path(pasta_limpa) / arquivo
        else:
            sqlite_path = default_sqlite

        if not sqlite_path.exists():
            sqlite_path = default_sqlite

        return create_engine(f"sqlite:///{sqlite_path.as_posix()}")

    return create_engine(f"sqlite:///{default_sqlite.as_posix()}")


def parse_sqls(caminho_sql: Path) -> list[dict]:
    conteudo = caminho_sql.read_text(encoding="utf-8")
    linhas = conteudo.splitlines()

    consultas = []
    bloco_atual = []
    comentarios = []
    titulo_atual = ""

    for linha in linhas:
        strip = linha.strip()

        if strip.startswith("#"):
            comentarios.append(strip.lstrip("#").strip())
            continue

        if re.match(r"^\s*SELECT\b", linha, flags=re.IGNORECASE):
            if bloco_atual:
                consultas.append({"titulo": titulo_atual, "sql": "\n".join(bloco_atual).strip()})
                bloco_atual = []

            titulo_atual = " | ".join([c for c in comentarios if c])
            if not titulo_atual:
                titulo_atual = f"Consulta {len(consultas) + 1}"
            comentarios = []

            bloco_atual.append(linha)
            continue

        if bloco_atual:
            bloco_atual.append(linha)

    if bloco_atual:
        consultas.append({"titulo": titulo_atual, "sql": "\n".join(bloco_atual).strip()})

    for i, q in enumerate(consultas, start=1):
        q["id"] = i
        if not q["titulo"]:
            q["titulo"] = f"Consulta {i}"

    return consultas


def aplicar_filtro_uf(sql: str, uf: str) -> tuple[str, dict]:
    if uf == "TODOS":
        return sql, {}

    if re.search(r"\bwhere\b", sql, flags=re.IGNORECASE):
        sql_filtrado = re.sub(
            r"\bwhere\b",
            "WHERE cc.estado = :uf AND ",
            sql,
            count=1,
            flags=re.IGNORECASE,
        )
        return sql_filtrado, {"uf": uf}

    if re.search(r"\bfrom\s+curso_censo\s+cc\b", sql, flags=re.IGNORECASE):
        sql_filtrado = f"{sql}\nWHERE cc.estado = :uf"
        return sql_filtrado, {"uf": uf}

    return sql, {}


def detectar_metricas_agregadas(sql: str, colunas_df: list[str]) -> dict[str, str]:
    metricas = {}
    padrao = re.compile(
        r"\b(count|sum)\s*\([^\)]*\)\s*(?:as\s+)?([a-zA-Z_][a-zA-Z0-9_]*)",
        flags=re.IGNORECASE,
    )

    for match in padrao.finditer(sql):
        agg = match.group(1).lower()
        alias = match.group(2).lower()
        metricas[alias] = agg

    resposta = {}
    for col in colunas_df:
        chave = col.lower()
        resposta[col] = metricas.get(chave, "agg")

    return resposta


def sanitizar_nome_aba(nome: str, existentes: set[str]) -> str:
    proibidos = r"[\\/*?:\[\]]"
    limpo = re.sub(proibidos, " ", nome).strip()
    limpo = re.sub(r"\s+", " ", limpo)
    if not limpo:
        limpo = "Planilha"
    limpo = limpo[:31]

    base = limpo
    i = 2
    while limpo in existentes:
        sufixo = f"_{i}"
        limpo = f"{base[:31-len(sufixo)]}{sufixo}"
        i += 1

    existentes.add(limpo)
    return limpo


def montar_tabela_metrica(df: pd.DataFrame, metrica: str) -> pd.DataFrame:
    base = df.copy()
    if "ano_censo" not in base.columns:
        base["ano_censo"] = "sem_ano"

    base["ano_censo"] = pd.to_numeric(base["ano_censo"], errors="coerce")
    base = base.dropna(subset=["ano_censo"])
    base["ano_censo"] = base["ano_censo"].astype(int)

    colunas_dim = [c for c in base.columns if c not in ["ano_censo", metrica]]

    if not colunas_dim:
        base["categoria"] = "geral"
        colunas_dim = ["categoria"]

    agrupado = (
        base.groupby(colunas_dim + ["ano_censo"], dropna=False, as_index=False)[metrica]
        .sum()
    )

    tabela = agrupado.pivot_table(
        index=colunas_dim,
        columns="ano_censo",
        values=metrica,
        aggfunc="sum",
        fill_value=0,
    ).reset_index()

    col_anos = sorted([c for c in tabela.columns if isinstance(c, int)])
    colunas_finais = colunas_dim + col_anos
    tabela = tabela[colunas_finais]
    return tabela


def estilo_moderno_excel(caminho_xlsx: Path):
    wb = load_workbook(caminho_xlsx)

    cor_header = PatternFill("solid", fgColor="0F172A")
    cor_subheader = PatternFill("solid", fgColor="1D4ED8")
    cor_zebra = PatternFill("solid", fgColor="EEF2FF")
    cor_branca = PatternFill("solid", fgColor="FFFFFF")
    fonte_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fonte_padrao = Font(name="Calibri", size=10, color="111827")
    borda_fina = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        for cell in ws[1]:
            cell.fill = cor_header
            cell.font = fonte_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = borda_fina

        max_col = ws.max_column
        max_row = ws.max_row

        for col_idx in range(1, max_col + 1):
            letra = ws.cell(row=1, column=col_idx).column_letter
            if col_idx <= 3:
                ws.column_dimensions[letra].width = 26
            else:
                ws.column_dimensions[letra].width = 14

        for row_idx in range(2, max_row + 1):
            fill = cor_zebra if row_idx % 2 == 0 else cor_branca
            for col_idx in range(1, max_col + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.fill = fill
                c.font = fonte_padrao
                c.border = borda_fina
                if isinstance(c.value, (int, float)):
                    c.number_format = "#,##0"

        for col_idx in range(1, min(3, max_col) + 1):
            ws.cell(row=1, column=col_idx).fill = cor_subheader

        ws.auto_filter.ref = ws.dimensions

    wb.save(caminho_xlsx)


def solicitar_uf() -> str:
    print("\nInforme a UF desejada ou TODOS para Brasil inteiro.")
    print("Exemplos: SC, SP, DF, TODOS")

    while True:
        uf = input("UF: ").strip().upper()
        if uf == "TODOS" or uf in UF_SIGLAS:
            return uf
        print("Valor inválido. Digite uma UF válida (AC..TO) ou TODOS.")


def main():
    base_dir = Path(__file__).resolve().parent
    caminho_sql = base_dir / "scripts.sql"
    consultas = parse_sqls(caminho_sql)
    uf = solicitar_uf()

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = base_dir / f"licenciaturas_sql_{uf.lower()}_{timestamp}.xlsx"

    engine = carregar_engine()
    nomes_abas = set()
    total_abas = 0

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for consulta in consultas:
            sql = consulta["sql"]
            titulo = consulta["titulo"]
            sql_exec, params = aplicar_filtro_uf(sql, uf)

            try:
                df = pd.read_sql(text(sql_exec), engine, params=params)
            except Exception as exc:
                print(f"[ERRO] Consulta {consulta['id']:02d} falhou: {exc}")
                continue

            if df.empty:
                print(f"[AVISO] Consulta {consulta['id']:02d} sem dados.")
                continue

            if "ano_censo" not in df.columns:
                print(f"[AVISO] Consulta {consulta['id']:02d} sem coluna ano_censo. Ignorada.")
                continue

            metricas = [
                c for c in df.columns
                if c != "ano_censo" and pd.api.types.is_numeric_dtype(df[c])
            ]
            if not metricas:
                print(f"[AVISO] Consulta {consulta['id']:02d} sem colunas numéricas agregadas. Ignorada.")
                continue

            tipo_metricas = detectar_metricas_agregadas(sql, metricas)

            for metrica in metricas:
                try:
                    tabela = montar_tabela_metrica(df, metrica)
                except Exception as exc:
                    print(
                        f"[AVISO] Métrica {metrica} da consulta {consulta['id']:02d} não pôde ser pivotada: {exc}"
                    )
                    continue

                tipo = tipo_metricas.get(metrica, "agg")
                nome_aba_bruto = f"Q{consulta['id']:02d}_{tipo}_{metrica}"
                nome_aba = sanitizar_nome_aba(nome_aba_bruto, nomes_abas)

                tabela.to_excel(writer, sheet_name=nome_aba, index=False)
                total_abas += 1

                print(
                    f"[OK] Consulta {consulta['id']:02d} | {titulo[:60]} | "
                    f"métrica={metrica} -> aba={nome_aba}"
                )

    estilo_moderno_excel(out)

    print("\n----------------------------------------------")
    print("Planilha gerada com sucesso")
    print(f"Arquivo: {out}")
    print(f"UF: {uf}")
    print(f"Total de abas: {total_abas}")
    print("----------------------------------------------")


if __name__ == "__main__":
    main()
