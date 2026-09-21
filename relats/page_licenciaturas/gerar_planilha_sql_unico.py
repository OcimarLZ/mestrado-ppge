import argparse
import configparser
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine, text


UF_SIGLAS = [
    "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT",
    "MS", "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO",
    "RR", "SC", "SP", "SE", "TO",
]


def carregar_engine() -> object:
    root = Path(__file__).resolve().parents[2]
    config_path = root / "config.ini"
    default_sqlite = root / "INEP.db"

    if not config_path.exists():
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    parser = configparser.ConfigParser()
    parser.read(config_path, encoding="utf-8")

    if "BD_INEP" not in parser:
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    bd = parser["BD_INEP"]
    sgdb = bd.get("SGDB", "sqlite").strip().lower()

    if sgdb == "sqlite":
        pasta_local = bd.get("PASTA_LOCAL", "").replace("\\", "/")
        arquivo = bd.get("DW", "INEP.db").strip()
        if pasta_local:
            sqlite_path = Path(pasta_local.strip("/")) / arquivo
        else:
            sqlite_path = default_sqlite

        if not sqlite_path.exists():
            sqlite_path = default_sqlite

        return create_engine(f"sqlite:///{sqlite_path.as_posix()}")

    return create_engine(f"sqlite:///{default_sqlite.as_posix()}")


def escolher_coluna_ano(colunas: list[str]) -> str:
    print("\nColunas retornadas pelo SQL:")
    for i, c in enumerate(colunas, start=1):
        print(f"{i:02d} - {c}")

    sugestao = 1
    while True:
        entrada = input(
            f"\nQual coluna representa os anos? [Enter para {sugestao:02d}] "
        ).strip()
        if not entrada:
            return colunas[sugestao - 1]
        if entrada.isdigit():
            idx = int(entrada)
            if 1 <= idx <= len(colunas):
                return colunas[idx - 1]
        if entrada in colunas:
            return entrada
        print("Valor inválido. Informe o número da coluna ou o nome exato.")


def escolher_colunas_legenda(colunas: list[str], coluna_ano: str) -> list[str]:
    elegiveis = [c for c in colunas if c != coluna_ano]

    print("\nQuais colunas serão de legenda/categoria (linhas)?")
    print("Digite índices separados por vírgula (ex: 1,3).")
    print("Se deixar vazio, será usado agrupamento geral.")

    for i, c in enumerate(elegiveis, start=1):
        print(f"{i:02d} - {c}")

    while True:
        entrada = input("Colunas de legenda: ").strip()
        if not entrada:
            return []

        tokens = [t.strip() for t in entrada.split(",") if t.strip()]
        ok = True
        escolhas = []

        for token in tokens:
            if token.isdigit():
                idx = int(token)
                if 1 <= idx <= len(elegiveis):
                    escolhas.append(elegiveis[idx - 1])
                else:
                    ok = False
                    break
            elif token in elegiveis:
                escolhas.append(token)
            else:
                ok = False
                break

        if ok:
            # remove duplicadas preservando ordem
            return list(dict.fromkeys(escolhas))

        print("Entrada inválida. Tente novamente.")


def escolher_mesclar_legendas() -> bool:
    while True:
        entrada = input("Mesclar colunas de legenda em uma só? [S/n] ").strip().lower()
        if entrada in ["", "s", "sim", "y", "yes"]:
            return True
        if entrada in ["n", "nao", "não"]:
            return False
        print("Resposta inválida. Use s ou n.")


def escolher_rotulo_dados() -> bool:
    while True:
        entrada = input("Inserir rótulo dos dados no gráfico? [s/N] ").strip().lower()
        if entrada in ["s", "sim", "y", "yes"]:
            return True
        if entrada in ["", "n", "nao", "não"]:
            return False
        print("Resposta inválida. Use s ou n.")


def escolher_uf() -> str:
    print("\nFiltro por UF:")
    print("0 - TODOS")
    for i, uf in enumerate(UF_SIGLAS, start=1):
        print(f"{i:02d} - {uf}")

    while True:
        entrada = input("Digite a opção (número) ou a sigla da UF: ").strip().upper()
        if entrada in ["0", "TODOS"]:
            return "TODOS"

        if entrada in UF_SIGLAS:
            return entrada

        if entrada.isdigit():
            idx = int(entrada)
            if 1 <= idx <= len(UF_SIGLAS):
                return UF_SIGLAS[idx - 1]

        print("Opção inválida. Tente novamente.")


def aplicar_filtro_uf(sql: str, uf: str) -> tuple[str, dict]:
    if uf == "TODOS":
        return sql, {}

    if re.search(r"\bfrom\s+curso_censo\s+cc\b", sql, flags=re.IGNORECASE):
        filtro = "cc.estado = :uf"
    elif re.search(r"\bfrom\s+curso_censo\b", sql, flags=re.IGNORECASE):
        filtro = "curso_censo.estado = :uf"
    else:
        print("[AVISO] SQL não usa curso_censo diretamente; filtro por UF não foi aplicado.")
        return sql, {}

    if re.search(r"\bwhere\b", sql, flags=re.IGNORECASE):
        sql_filtrado = re.sub(
            r"\bwhere\b",
            f"WHERE {filtro} AND ",
            sql,
            count=1,
            flags=re.IGNORECASE,
        )
        return sql_filtrado, {"uf": uf}

    sql_filtrado = re.sub(
        r"\b(group\s+by|order\s+by|limit)\b",
        f" WHERE {filtro} \\1",
        sql,
        count=1,
        flags=re.IGNORECASE,
    )
    if sql_filtrado == sql:
        sql_filtrado = f"{sql}\nWHERE {filtro}"

    return sql_filtrado, {"uf": uf}


def aplicar_title_case_legendas(df: pd.DataFrame, colunas_legenda: list[str]) -> None:
    """Aplica title case aos valores das colunas de legenda."""
    for coluna in colunas_legenda:
        if coluna in df.columns:
            df[coluna] = df[coluna].astype(str).str.title()


def montar_planilha(
    df: pd.DataFrame,
    coluna_ano: str,
    colunas_legenda: list[str],
    mesclar_legendas: bool,
) -> pd.DataFrame:
    base = df.copy()
    base[coluna_ano] = pd.to_numeric(base[coluna_ano], errors="coerce")
    base = base.dropna(subset=[coluna_ano])

    if base.empty:
        return base

    base[coluna_ano] = base[coluna_ano].astype(int)

    colunas_valor = [
        c for c in base.columns
        if c != coluna_ano and c not in colunas_legenda and pd.api.types.is_numeric_dtype(base[c])
    ]

    if not colunas_valor:
        restantes = [c for c in base.columns if c != coluna_ano and c not in colunas_legenda]
        if restantes:
            colunas_valor = [restantes[0]]
            base[colunas_valor[0]] = 1
        else:
            base["valor"] = 1
            colunas_valor = ["valor"]

    if not colunas_legenda:
        base["categoria"] = "Geral"
        colunas_legenda_uso = ["categoria"]
    else:
        colunas_legenda_uso = colunas_legenda.copy()
        aplicar_title_case_legendas(base, colunas_legenda_uso)

    if mesclar_legendas and len(colunas_legenda_uso) > 1:
        base["legenda"] = base[colunas_legenda_uso].astype(str).agg(" | ".join, axis=1)
        colunas_legenda_uso = ["legenda"]

    if len(colunas_valor) > 1:
        longo = base.melt(
            id_vars=colunas_legenda_uso + [coluna_ano],
            value_vars=colunas_valor,
            var_name="agrupamento",
            value_name="valor",
        )
        # Aplica title case às colunas de legenda após melt
        aplicar_title_case_legendas(longo, colunas_legenda_uso)
        if mesclar_legendas:
            cols_merge = colunas_legenda_uso + ["agrupamento"]
            longo["legenda"] = longo[cols_merge].astype(str).agg(" | ".join, axis=1)
            indice = ["legenda"]
        else:
            indice = colunas_legenda_uso + ["agrupamento"]
        tabela = longo.pivot_table(
            index=indice,
            columns=coluna_ano,
            values="valor",
            aggfunc="sum",
            fill_value=0,
        ).reset_index()
    else:
        valor_col = colunas_valor[0]
        tabela = base.pivot_table(
            index=colunas_legenda_uso,
            columns=coluna_ano,
            values=valor_col,
            aggfunc="sum",
            fill_value=0,
        ).reset_index()

    colunas_anos = sorted([c for c in tabela.columns if isinstance(c, int)])
    colunas_fixas = [c for c in tabela.columns if c not in colunas_anos]
    tabela = tabela[colunas_fixas + colunas_anos]

    # Mantem anos como texto no cabecalho para o eixo X do grafico aparecer corretamente.
    mapa_anos = {c: str(c) for c in colunas_anos}
    tabela = tabela.rename(columns=mapa_anos)
    return tabela


def adicionar_grafico_linhas(ws, tabela: pd.DataFrame, mostrar_rotulo_dados: bool = False):
    if tabela.empty:
        return

    colunas = list(tabela.columns)
    colunas_anos = [
        c for c in colunas
        if isinstance(c, int) or (isinstance(c, str) and c.isdigit())
    ]
    if not colunas_anos:
        return

    idx_primeiro_ano = colunas.index(colunas_anos[0]) + 1
    idx_ultimo_ano = colunas.index(colunas_anos[-1]) + 1

    chart = LineChart()
    chart.style = 10
    chart.height = 11
    chart.width = 22
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.tickLblPos = "low"
    chart.y_axis.number_format = "#,##0"
    chart.legend.position = "b"
    chart.legend.overlay = False

    max_series = min(len(tabela), 20)
    if max_series <= 0:
        return

    bloco_valores = Reference(
        ws,
        min_col=idx_primeiro_ano,
        max_col=idx_ultimo_ano,
        min_row=2,
        max_row=1 + max_series,
    )
    chart.add_data(bloco_valores, titles_from_data=False, from_rows=True)

    categorias = Reference(
        ws,
        min_col=idx_primeiro_ano,
        max_col=idx_ultimo_ano,
        min_row=1,
        max_row=1,
    )
    chart.set_categories(categorias)

    if mostrar_rotulo_dados:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.dLblPos = "t"
        chart.dataLabels.showSerName = False
        chart.dataLabels.showCatName = False
        chart.dataLabels.showLegendKey = False
        chart.dataLabels.showPercent = False
        chart.dataLabels.showBubbleSize = False

    for i, linha in enumerate(range(2, 2 + max_series)):
        partes = []
        for col in range(1, idx_primeiro_ano):
            valor = ws.cell(row=linha, column=col).value
            if valor is not None and str(valor).strip() != "":
                partes.append(str(valor))
        titulo = " | ".join(partes) if partes else f"serie_{linha - 1}"

        if i < len(chart.series):
            try:
                chart.series[i].title = SeriesLabel(v=titulo)
            except Exception:
                try:
                    chart.series[i].title = titulo
                except Exception:
                    pass

    ancora = f"{get_column_letter(max(1, idx_primeiro_ano))}{len(tabela) + 4}"
    ws.add_chart(chart, ancora)


def aplicar_formatacao_tabela(ws, tabela: pd.DataFrame):
    fonte = Font(name="Times New Roman", size=10)

    max_row = ws.max_row
    max_col = ws.max_column

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            celula = ws.cell(row=row, column=col)
            celula.font = fonte

            if row >= 2 and isinstance(celula.value, (int, float)):
                celula.number_format = "#,##0"

    # Ajuste simples de largura com base no conteúdo
    for col in range(1, max_col + 1):
        letra = get_column_letter(col)
        tamanho = 10
        for row in range(1, min(max_row, 1000) + 1):
            valor = ws.cell(row=row, column=col).value
            if valor is not None:
                tamanho = max(tamanho, len(str(valor)) + 2)
        ws.column_dimensions[letra].width = min(tamanho, 60)


def ler_sql(caminho_sql: Path) -> str:
    conteudo = caminho_sql.read_text(encoding="utf-8").strip()
    if not conteudo:
        raise ValueError("Arquivo SQL está vazio.")
    return conteudo


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Executa um único arquivo .sql e gera um .xlsx com anos em colunas e "
            "legendas em linhas."
        )
    )
    parser.add_argument(
        "sql_file",
        help="Caminho para o arquivo .sql a ser executado.",
    )
    return parser.parse_args()


def salvar_excel_com_fallback(saida: Path, tabela: pd.DataFrame, mostrar_rotulo_dados: bool = False) -> Path:
    try:
        with pd.ExcelWriter(saida, engine="openpyxl") as writer:
            tabela.to_excel(writer, sheet_name="dados", index=False)
            ws = writer.book["dados"]
            aplicar_formatacao_tabela(ws, tabela)
            adicionar_grafico_linhas(ws, tabela, mostrar_rotulo_dados=mostrar_rotulo_dados)
        return saida
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        saida_alt = saida.with_name(f"{saida.stem}_{timestamp}{saida.suffix}")
        print(
            "[AVISO] O arquivo de saída estava em uso/bloqueado. "
            f"Salvando em arquivo alternativo: {saida_alt.name}"
        )
        with pd.ExcelWriter(saida_alt, engine="openpyxl") as writer:
            tabela.to_excel(writer, sheet_name="dados", index=False)
            ws = writer.book["dados"]
            aplicar_formatacao_tabela(ws, tabela)
            adicionar_grafico_linhas(ws, tabela, mostrar_rotulo_dados=mostrar_rotulo_dados)
        return saida_alt


def main():
    args = parse_args()
    caminho_sql = Path(args.sql_file).resolve()

    if not caminho_sql.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho_sql}")
    if caminho_sql.suffix.lower() != ".sql":
        raise ValueError("O parâmetro deve apontar para um arquivo com extensão .sql")

    sql = ler_sql(caminho_sql)
    uf = escolher_uf()
    sql_exec, params = aplicar_filtro_uf(sql, uf)
    engine = carregar_engine()

    df = pd.read_sql(text(sql_exec), engine, params=params)
    if df.empty:
        print("Consulta sem dados. Nenhuma planilha foi gerada.")
        return

    coluna_ano = escolher_coluna_ano(list(df.columns))
    colunas_legenda = escolher_colunas_legenda(list(df.columns), coluna_ano)
    mesclar_legendas = escolher_mesclar_legendas() if colunas_legenda else False
    mostrar_rotulo_dados = escolher_rotulo_dados()

    tabela = montar_planilha(df, coluna_ano, colunas_legenda, mesclar_legendas)

    if uf == "TODOS":
        saida = caminho_sql.with_suffix(".xlsx")
    else:
        saida = caminho_sql.with_name(f"{caminho_sql.stem}_{uf}{caminho_sql.suffix}").with_suffix(".xlsx")
    saida_real = salvar_excel_com_fallback(saida, tabela, mostrar_rotulo_dados=mostrar_rotulo_dados)

    print("\nPlanilha gerada com sucesso.")
    print(f"Arquivo: {saida_real}")
    print(f"Filtro de UF: {uf}")
    print(f"Coluna de ano: {coluna_ano}")
    print(f"Colunas de legenda: {colunas_legenda if colunas_legenda else ['geral']}")
    print(f"Legendas mescladas: {'sim' if mesclar_legendas else 'não'}")
    print(f"Rótulo dos dados no gráfico: {'sim' if mostrar_rotulo_dados else 'não'}")


if __name__ == "__main__":
    main()
