import argparse
import configparser
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from sqlalchemy import create_engine, text

def carregar_engine() -> object:
    root = Path(__file__).resolve().parents[2]
    config_path = root / "config.ini"
    default_sqlite = root / "INEP.db"

    if not config_path.exists():
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    parser = configparser.ConfigParser()
    parser.read(config_path, encoding="utf-8")

    if "BD_INEP" not in parser:
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    bd = parser["BD_INEP"]
    sgdb = bd.get("SGDB", "sqlite").strip().lower()

    if sgdb == "sqlite":
        pasta_local = bd.get("PASTA_LOCAL", "").replace("\\", "/")
        arquivo = bd.get("DW", "INEP.db").strip()
        if pasta_local:
            sqlite_path = Path(pasta_local.strip("/")) / arquivo
        else:
            sqlite_path = default_sqlite

        if not sqlite_path.exists():
            sqlite_path = default_sqlite

        return create_engine(f"sqlite:///{sqlite_path.as_posix()}")

    return create_engine(f"sqlite:///{default_sqlite.as_posix()}")


def main():
    parser = argparse.ArgumentParser(description="Gera uma planilha resumo (Cursos, Campus, Polos) por IES.")
    parser.add_argument("sigla", help="Sigla da IES para filtrar (ex: UFFS)")
    args = parser.parse_args()

    sigla = args.sigla.upper()

    sql = """
    SELECT distinct
    i.nome as ies,
    cc.ano_censo, 
    c.nome as curso,
    m.nome || '/' || m.estado as municipio,
    cc.tp_modalidade_ensino
    from curso_censo cc 
    join curso c on c.codigo = cc.curso 
    join ies i on i.codigo = cc.ies 
    join municipio m on m.codigo = cc.municipio  
    where cc.ano_censo in (2014, 2024) and cc.tp_grau_academico = 2 and i.sigla = :sigla
    """

    engine = carregar_engine()
    print(f"Executando a consulta para a sigla: {sigla}...")
    
    # 1. Obter código e nome da IES
    with engine.connect() as conn:
        res_ies = conn.execute(
            text("SELECT codigo, nome FROM ies WHERE sigla = :sigla"),
            {"sigla": sigla}
        ).fetchone()
        
    if not res_ies:
        print(f"Nenhum dado encontrado para a sigla {sigla}.")
        return
        
    ies_codigo, ies_nome = res_ies

    # 2. Consultar métricas de resumo
    def obter_metricas(ies_id, ano):
        with engine.connect() as conn:
            # a) Total de matriculas geral
            tot_mat = conn.execute(text("""
                SELECT SUM(qt_mat) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0
            
            # b) Total de matriculas em Licenciaturas
            tot_mat_lic = conn.execute(text("""
                SELECT SUM(qt_mat) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0
            
            # c) Numero de cursos em Licenciaturas
            num_cursos_lic = conn.execute(text("""
                SELECT COUNT(DISTINCT curso) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0
            
            # d) Numero de campus em Licenciaturas
            num_campus_lic = conn.execute(text("""
                SELECT COUNT(DISTINCT municipio) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2 AND tp_modalidade_ensino = 1
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0
            
            # d2) Numero de polos EaD em Licenciaturas
            num_polos_ead_lic = conn.execute(text("""
                SELECT COUNT(DISTINCT municipio) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2 AND tp_modalidade_ensino = 2
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0
            
            # e) Numero de polos em Licenciaturas
            num_polos_lic = conn.execute(text("""
                SELECT COUNT(DISTINCT id_polo) FROM uab_censo 
                WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0
            
        return tot_mat, tot_mat_lic, num_cursos_lic, num_campus_lic, num_polos_ead_lic, num_polos_lic

    tot_mat_2014, tot_mat_lic_2014, num_cursos_lic_2014, num_campus_lic_2014, num_polos_ead_lic_2014, num_polos_lic_2014 = obter_metricas(ies_codigo, 2014)
    tot_mat_2024, tot_mat_lic_2024, num_cursos_lic_2024, num_campus_lic_2024, num_polos_ead_lic_2024, num_polos_lic_2024 = obter_metricas(ies_codigo, 2024)

    sql_polos = """
    WITH ead AS (
        SELECT DISTINCT cc.ano_censo, m.nome || '/' || m.estado as municipio
        FROM curso_censo cc 
        JOIN ies i ON i.codigo = cc.ies 
        JOIN municipio m ON m.codigo = cc.municipio  
        WHERE cc.ano_censo IN (2014, 2024) AND cc.tp_grau_academico = 2 AND cc.tp_modalidade_ensino = 2 AND i.sigla = :sigla
    ), uab AS (
        SELECT DISTINCT uc.ano_censo, m.nome || '/' || m.estado as municipio
        FROM uab_censo uc 
        JOIN ies i ON i.codigo = uc.ies 
        JOIN municipio m ON m.codigo = uc.municipio  
        WHERE uc.ano_censo IN (2014, 2024) AND uc.tp_grau_academico = 2 AND i.sigla = :sigla
    ), combinados AS (
        SELECT ano_censo, municipio FROM ead
        UNION
        SELECT ano_censo, municipio FROM uab
    )
    SELECT c.ano_censo, 
           CASE WHEN u.municipio IS NOT NULL THEN c.municipio || ' (UAB)' ELSE c.municipio END as municipio
    FROM combinados c
    LEFT JOIN uab u ON c.ano_censo = u.ano_censo AND c.municipio = u.municipio
    """

    df = pd.read_sql(text(sql), engine, params={"sigla": sigla})
    df_polos = pd.read_sql(text(sql_polos), engine, params={"sigla": sigla})
    
    if df.empty:
        print(f"Nenhum dado detalhado encontrado para a sigla {sigla} em curso_censo.")
        return

    # Processamento para 2014
    df_2014 = df[df['ano_censo'] == 2014]
    cursos_2014 = sorted(df_2014['curso'].dropna().unique().tolist())
    campus_2014 = sorted(df_2014[df_2014['tp_modalidade_ensino'] == 1]['municipio'].dropna().unique().tolist())
    
    df_polos_2014 = df_polos[df_polos['ano_censo'] == 2014]
    polos_2014 = sorted(df_polos_2014['municipio'].dropna().unique().tolist())

    # Processamento para 2024
    df_2024 = df[df['ano_censo'] == 2024]
    cursos_2024 = sorted(df_2024['curso'].dropna().unique().tolist())
    campus_2024 = sorted(df_2024[df_2024['tp_modalidade_ensino'] == 1]['municipio'].dropna().unique().tolist())
    
    df_polos_2024 = df_polos[df_polos['ano_censo'] == 2024]
    polos_2024 = sorted(df_polos_2024['municipio'].dropna().unique().tolist())

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"

    # a) Coluna "ies" mesclada (agora em 6 colunas, A a F)
    ws.merge_cells('A1:F1')
    ws['A1'] = ies_nome
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

    # b) Anos 2014 (A a C) e 2024 (D a F) - Agora na linha 2
    ws.merge_cells('A2:C2')
    ws['A2'] = '2014'
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('D2:F2')
    ws['D2'] = '2024'
    ws['D2'].alignment = Alignment(horizontal="center", vertical="center")

    # Inserção das novas linhas de métricas
    # Linha 3: Total de matrículas
    ws['A3'] = "Total de matrículas no ano"
    ws['B3'] = tot_mat_2014
    ws['D3'] = "Total de matrículas no ano"
    ws['E3'] = tot_mat_2024
    
    # Linha 4: Total de matrículas em Licenciaturas
    ws['A4'] = "Matrículas em Licenciaturas"
    ws['B4'] = tot_mat_lic_2014
    ws['D4'] = "Matrículas em Licenciaturas"
    ws['E4'] = tot_mat_lic_2024
    
    # Linha 5: Número de cursos de Licenciatura
    ws['A5'] = "Número de cursos de Licenciatura"
    ws['B5'] = num_cursos_lic_2014
    ws['D5'] = "Número de cursos de Licenciatura"
    ws['E5'] = num_cursos_lic_2024
    
    # Linha 6: Número de campus (municípios)
    ws['A6'] = "Número de campus (municípios)"
    ws['B6'] = num_campus_lic_2014
    ws['D6'] = "Número de campus (municípios)"
    ws['E6'] = num_campus_lic_2024
    
    # Linha 7: Número de polos EaD
    ws['A7'] = "Número de polos EaD"
    ws['B7'] = num_polos_ead_lic_2014
    ws['D7'] = "Número de polos EaD"
    ws['E7'] = num_polos_ead_lic_2024
    
    # Linha 8: Número de polos (UAB)
    ws['A8'] = "Número de polos (UAB)"
    ws['B8'] = num_polos_lic_2014
    ws['D8'] = "Número de polos (UAB)"
    ws['E8'] = num_polos_lic_2024

    # Mesclar colunas B/C e E/F das linhas 3 a 8 para as métricas
    for r in range(3, 9):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)

    # Cabeçalhos - Agora na linha 9
    headers = ['Cursos', 'Campus', 'Polos EaD/UAB', 'Cursos', 'Campus', 'Polos EaD/UAB']
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=9, column=col_idx, value=header)

    # Preenchendo os dados (Listas de tamanhos diferentes) - Agora a partir da linha 10
    max_rows = max(len(cursos_2014), len(campus_2014), len(polos_2014), 
                   len(cursos_2024), len(campus_2024), len(polos_2024))

    for i in range(max_rows):
        row_idx = i + 10
        
        # 2014
        if i < len(cursos_2014):
            ws.cell(row=row_idx, column=1, value=cursos_2014[i])
        if i < len(campus_2014):
            ws.cell(row=row_idx, column=2, value=campus_2014[i])
        if i < len(polos_2014):
            ws.cell(row=row_idx, column=3, value=polos_2014[i])
            
        # 2024
        if i < len(cursos_2024):
            ws.cell(row=row_idx, column=4, value=cursos_2024[i])
        if i < len(campus_2024):
            ws.cell(row=row_idx, column=5, value=campus_2024[i])
        if i < len(polos_2024):
            ws.cell(row=row_idx, column=6, value=polos_2024[i])

    # Estilos
    font_default = Font(name='Times New Roman', size=10)
    font_bold = Font(name='Times New Roman', size=10, bold=True)
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_green = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fill_orange = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    fill_grey = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    # Aplicamos a borda no retângulo inteiro dos dados para ficar alinhado
    for row in ws.iter_rows(min_row=1, max_row=max_rows+9, min_col=1, max_col=6):
        for cell in row:
            cell.border = thin_border
            if cell.row == 1:
                cell.font = font_bold
                cell.fill = fill_yellow
            elif cell.row == 2:
                cell.font = font_bold
                if cell.column <= 3:
                    cell.fill = fill_green
                else:
                    cell.fill = fill_orange
            elif 3 <= cell.row <= 8:
                if cell.column in (1, 4):
                    cell.font = font_default
                elif cell.column in (2, 5):
                    cell.font = font_bold
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.font = font_default
            elif cell.row == 9:
                cell.font = font_bold
                cell.fill = fill_grey
            else:
                cell.font = font_default

    # Largura das colunas
    for col in ['A', 'D']:
        ws.column_dimensions[col].width = 35 # Cursos
    for col in ['B', 'E', 'C', 'F']:
        ws.column_dimensions[col].width = 25 # Campus e Polos

    arquivo_saida = Path(__file__).parent / f"resumo_ies_{sigla}.xlsx"
    wb.save(arquivo_saida)
    
    print(f"Planilha gerada com sucesso: {arquivo_saida}")

if __name__ == "__main__":
    main()
