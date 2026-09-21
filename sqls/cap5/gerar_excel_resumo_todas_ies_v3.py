import configparser
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from sqlalchemy import create_engine, text


def carregar_engine():
    root = Path(__file__).resolve().parents[2]
    config_path = root / "config.ini"
    default_sqlite = root / "INEP.db"

    if not config_path.exists():
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    parser = configparser.ConfigParser()
    parser.read(config_path, encoding="utf-8")

    if "BD_INEP" not in parser:
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    bd = parser["BD_INEP"]
    sgdb = bd.get("SGDB", "sqlite").strip().lower()

    if sgdb == "sqlite":
        pasta_local = bd.get("PASTA_LOCAL", "").replace("\\", "/")
        arquivo = bd.get("DW", "INEP.db").strip()
        if pasta_local:
            sqlite_path = Path(pasta_local.strip("/")) / arquivo
        else:
            sqlite_path = default_sqlite

        if not sqlite_path.exists():
            sqlite_path = default_sqlite

        return create_engine(f"sqlite:///{sqlite_path.as_posix()}")

    return create_engine(f"sqlite:///{default_sqlite.as_posix()}")


def processar_ies(engine, sigla):
    sql = """
    SELECT distinct
    i.nome as ies,
    cc.ano_censo, 
    c.nome || ' (' || c.codigo || ')' as curso,
    m.nome || '/' || m.estado as municipio,
    cc.tp_modalidade_ensino
    from curso_censo cc 
    join curso c on c.codigo = cc.curso 
    join ies i on i.codigo = cc.ies 
    join municipio m on m.codigo = cc.municipio  
    where cc.ano_censo in (2014, 2024) and cc.tp_grau_academico = 2 and i.sigla = :sigla
    """

    with engine.connect() as conn:
        res_ies = conn.execute(
            text("SELECT codigo, nome FROM ies WHERE sigla = :sigla"),
            {"sigla": sigla}
        ).fetchone()

    if not res_ies:
        return None

    ies_codigo, ies_nome = res_ies

    def obter_metricas(ies_id, ano):
        with engine.connect() as conn:
            tot_mat = conn.execute(text("""
                SELECT SUM(qt_mat) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0

            tot_mat_lic = conn.execute(text("""
                SELECT SUM(qt_mat) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0

            num_cursos_lic = conn.execute(text("""
                SELECT COUNT(DISTINCT curso) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0

            num_campus_lic = conn.execute(text("""
                SELECT COUNT(DISTINCT municipio) FROM curso_censo 
                WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2 AND tp_modalidade_ensino = 1
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0

            num_polos_ead_proprios = conn.execute(text("""
                SELECT COUNT(DISTINCT cc.municipio) 
                FROM curso_censo cc
                JOIN curso c ON c.codigo = cc.curso
                WHERE cc.ies = :ies AND cc.ano_censo = :ano AND cc.tp_grau_academico = 2 AND cc.tp_modalidade_ensino = 2 AND c.fl_uab = 'N'
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0

            num_polos_lic = conn.execute(text("""
                SELECT COUNT(DISTINCT municipio) FROM (
                    SELECT municipio FROM uab_censo 
                    WHERE ies = :ies AND ano_censo = :ano AND tp_grau_academico = 2
                    UNION
                    SELECT cc.municipio FROM curso_censo cc
                    JOIN curso c ON c.codigo = cc.curso
                    WHERE cc.ies = :ies AND cc.ano_censo = :ano AND cc.tp_grau_academico = 2 AND cc.tp_modalidade_ensino = 2 AND c.fl_uab = 'S'
                )
            """), {"ies": ies_id, "ano": ano}).fetchone()[0] or 0
            
            num_polos_ead_lic = num_polos_ead_proprios + num_polos_lic

        return tot_mat, tot_mat_lic, num_cursos_lic, num_campus_lic, num_polos_ead_lic, num_polos_lic, num_polos_ead_proprios

    tot_mat_2014, tot_mat_lic_2014, num_cursos_lic_2014, num_campus_lic_2014, num_polos_ead_lic_2014, num_polos_lic_2014, num_polos_ead_proprios_2014 = obter_metricas(ies_codigo, 2014)
    tot_mat_2024, tot_mat_lic_2024, num_cursos_lic_2024, num_campus_lic_2024, num_polos_ead_lic_2024, num_polos_lic_2024, num_polos_ead_proprios_2024 = obter_metricas(ies_codigo, 2024)

    sql_polos = """
    WITH ead_proprio AS (
        SELECT DISTINCT cc.ano_censo, m.nome || '/' || m.estado as municipio
        FROM curso_censo cc 
        JOIN curso c ON c.codigo = cc.curso
        JOIN ies i ON i.codigo = cc.ies 
        JOIN municipio m ON m.codigo = cc.municipio  
        WHERE cc.ano_censo IN (2014, 2024) AND cc.tp_grau_academico = 2 AND cc.tp_modalidade_ensino = 2 AND c.fl_uab = 'N' AND i.sigla = :sigla
    ), uab AS (
        SELECT DISTINCT u.ano_censo, m.nome || '/' || m.estado || ' (UAB)' as municipio
        FROM uab_censo u
        JOIN ies i ON i.codigo = u.ies
        JOIN municipio m ON m.codigo = u.municipio
        WHERE u.ano_censo IN (2014, 2024) AND u.tp_grau_academico = 2 AND i.sigla = :sigla
        UNION
        SELECT DISTINCT cc.ano_censo, m.nome || '/' || m.estado || ' (UAB)' as municipio
        FROM curso_censo cc 
        JOIN curso c ON c.codigo = cc.curso
        JOIN ies i ON i.codigo = cc.ies 
        JOIN municipio m ON m.codigo = cc.municipio  
        WHERE cc.ano_censo IN (2014, 2024) AND cc.tp_grau_academico = 2 AND cc.tp_modalidade_ensino = 2 AND c.fl_uab = 'S' AND i.sigla = :sigla
    )
    SELECT ano_censo, municipio FROM ead_proprio
    UNION
    SELECT ano_censo, municipio FROM uab
    ORDER BY municipio
    """

    df = pd.read_sql(text(sql), engine, params={"sigla": sigla})
    df_polos = pd.read_sql(text(sql_polos), engine, params={"sigla": sigla})

    if df.empty:
        return None

    df_2014 = df[df['ano_censo'] == 2014]
    cursos_2014 = sorted(df_2014['curso'].dropna().unique().tolist())
    campus_2014 = sorted(df_2014[df_2014['tp_modalidade_ensino'] == 1]['municipio'].dropna().unique().tolist())
    df_polos_2014 = df_polos[df_polos['ano_censo'] == 2014]
    polos_2014 = sorted(df_polos_2014['municipio'].dropna().unique().tolist())

    df_2024 = df[df['ano_censo'] == 2024]
    cursos_2024 = sorted(df_2024['curso'].dropna().unique().tolist())
    campus_2024 = sorted(df_2024[df_2024['tp_modalidade_ensino'] == 1]['municipio'].dropna().unique().tolist())
    df_polos_2024 = df_polos[df_polos['ano_censo'] == 2024]
    polos_2024 = sorted(df_polos_2024['municipio'].dropna().unique().tolist())

    return {
        'sigla': sigla,
        'nome': ies_nome,
        'tot_mat_2014': tot_mat_2014,
        'tot_mat_lic_2014': tot_mat_lic_2014,
        'num_cursos_lic_2014': num_cursos_lic_2014,
        'num_campus_lic_2014': num_campus_lic_2014,
        'num_polos_ead_lic_2014': num_polos_ead_lic_2014,
        'num_polos_ead_proprios_2014': num_polos_ead_proprios_2014,
        'num_polos_lic_2014': num_polos_lic_2014,
        'tot_mat_2024': tot_mat_2024,
        'tot_mat_lic_2024': tot_mat_lic_2024,
        'num_cursos_lic_2024': num_cursos_lic_2024,
        'num_campus_lic_2024': num_campus_lic_2024,
        'num_polos_ead_lic_2024': num_polos_ead_lic_2024,
        'num_polos_ead_proprios_2024': num_polos_ead_proprios_2024,
        'num_polos_lic_2024': num_polos_lic_2024,
        'cursos_2014': cursos_2014,
        'campus_2014': campus_2014,
        'polos_2014': polos_2014,
        'cursos_2024': cursos_2024,
        'campus_2024': campus_2024,
        'polos_2024': polos_2024
    }


def preencher_aba(ws, dados_ies, start_row):
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=6)
    ws.cell(row=start_row, column=1, value=dados_ies['nome']).alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=start_row+1, start_column=1, end_row=start_row+1, end_column=3)
    ws.cell(row=start_row+1, column=1, value='2014').alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=start_row+1, start_column=4, end_row=start_row+1, end_column=6)
    ws.cell(row=start_row+1, column=4, value='2024').alignment = Alignment(horizontal="center", vertical="center")

    labels = [
        ("Total de matrículas no ano", dados_ies['tot_mat_2014'], dados_ies['tot_mat_2024']),
        ("Matrículas em Licenciaturas", dados_ies['tot_mat_lic_2014'], dados_ies['tot_mat_lic_2024']),
        ("Número de cursos de Licenciatura", dados_ies['num_cursos_lic_2014'], dados_ies['num_cursos_lic_2024']),
        ("Número de campus (municípios)", dados_ies['num_campus_lic_2014'], dados_ies['num_campus_lic_2024']),
        ("Número de polos EaD", dados_ies['num_polos_ead_lic_2014'], dados_ies['num_polos_ead_lic_2024']),
        ("Número de polos EaD próprios", dados_ies['num_polos_ead_proprios_2014'], dados_ies['num_polos_ead_proprios_2024']),
        ("Número de polos (UAB)", dados_ies['num_polos_lic_2014'], dados_ies['num_polos_lic_2024'])
    ]

    for idx, (label, val2014, val2024) in enumerate(labels):
        r = start_row + 2 + idx
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=val2014)
        ws.cell(row=r, column=4, value=label)
        ws.cell(row=r, column=5, value=val2024)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)

    headers = ['Cursos', 'Campus', 'Polos EaD/UAB', 'Cursos', 'Campus', 'Polos EaD/UAB']
    header_row = start_row + 9
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col_idx, value=header)

    max_rows = max(len(dados_ies['cursos_2014']), len(dados_ies['campus_2014']), len(dados_ies['polos_2014']),
                   len(dados_ies['cursos_2024']), len(dados_ies['campus_2024']), len(dados_ies['polos_2024']))

    for i in range(max_rows):
        row_idx = header_row + 1 + i
        if i < len(dados_ies['cursos_2014']):
            ws.cell(row=row_idx, column=1, value=dados_ies['cursos_2014'][i])
        if i < len(dados_ies['campus_2014']):
            ws.cell(row=row_idx, column=2, value=dados_ies['campus_2014'][i])
        if i < len(dados_ies['polos_2014']):
            ws.cell(row=row_idx, column=3, value=dados_ies['polos_2014'][i])
        if i < len(dados_ies['cursos_2024']):
            ws.cell(row=row_idx, column=4, value=dados_ies['cursos_2024'][i])
        if i < len(dados_ies['campus_2024']):
            ws.cell(row=row_idx, column=5, value=dados_ies['campus_2024'][i])
        if i < len(dados_ies['polos_2024']):
            ws.cell(row=row_idx, column=6, value=dados_ies['polos_2024'][i])

    font_default = Font(name='Times New Roman', size=10)
    font_bold = Font(name='Times New Roman', size=10, bold=True)
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_green = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fill_orange = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    fill_grey = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    for row_idx in range(start_row, start_row + 10 + max_rows):
        for col_idx in range(1, 7):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            rel_row = row_idx - start_row + 1
            if rel_row == 1:
                cell.font = font_bold
                cell.fill = fill_yellow
            elif rel_row == 2:
                cell.font = font_bold
                if cell.column <= 3:
                    cell.fill = fill_green
                else:
                    cell.fill = fill_orange
            elif 3 <= rel_row <= 9:
                if cell.column in (1, 4):
                    cell.font = font_default
                elif cell.column in (2, 5):
                    cell.font = font_bold
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.font = font_default
            elif rel_row == 10:
                cell.font = font_bold
                cell.fill = fill_grey
            else:
                cell.font = font_default

    for col in ['A', 'D']:
        ws.column_dimensions[col].width = 35
    for col in ['B', 'E', 'C', 'F']:
        ws.column_dimensions[col].width = 25

    next_start_row = start_row + 10 + max_rows
    ws.merge_cells(start_row=next_start_row, start_column=1, end_row=next_start_row, end_column=6)

    return next_start_row + 1


def main():
    import sys
    print("Iniciando script...", flush=True)
    engine = carregar_engine()
    print("Conectado ao banco!", flush=True)

    print("Obtendo lista de IES...", flush=True)
    with engine.connect() as conn:
        siglas_result = conn.execute(text("""
            SELECT DISTINCT i.sigla, i.nome
            FROM ies i 
            JOIN curso_censo cc ON i.codigo = cc.ies
            WHERE i.categoria = 1 AND i.org_academica = 1 AND i.sigla <> '-'
            ORDER BY i.sigla
        """)).fetchall()
    print(f"Total de IES: {len(siglas_result)}", flush=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo Todas IES"
    
    start_row = 1

    for idx, (sigla, nome) in enumerate(siglas_result, start=1):
        print(f"[{idx}/{len(siglas_result)}] Processando {sigla}...", flush=True)
        dados = processar_ies(engine, sigla)
        if dados:
            print(f"  -> Dados encontrados!", flush=True)
            start_row = preencher_aba(ws, dados, start_row)
        else:
            print(f"  -> Nenhum dado para {sigla}", flush=True)

    arquivo = Path(__file__).parent / "resumo_todas_ies_v3.xlsx"
    print(f"Salvando em {arquivo}...", flush=True)
    wb.save(arquivo)
    print("Feito!", flush=True)


if __name__ == "__main__":
    main()
