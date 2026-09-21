import argparse
import configparser
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from sqlalchemy import create_engine, text

def carregar_engine() -> object:
    # Assumindo que o script esteja em d:\ProjetosPY\inep\sqls\cap5\
    root = Path(__file__).resolve().parents[2]
    config_path = root / "config.ini"
    default_sqlite = root / "INEP.db"

    if not config_path.exists():
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    parser = configparser.ConfigParser()
    parser.read(config_path, encoding="utf-8")

    if "BD_INEP" not in parser:
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    bd = parser["BD_INEP"]
    sgdb = bd.get("SGDB", "sqlite").strip().lower()

    if sgdb == "sqlite":
        pasta_local = bd.get("PASTA_LOCAL", "").replace("\\", "/")
        arquivo = bd.get("DW", "INEP.db").strip()
        if pasta_local:
            sqlite_path = Path(pasta_local.strip("/")) / arquivo
        else:
            sqlite_path = default_sqlite

        if not sqlite_path.exists():
            sqlite_path = default_sqlite

        return create_engine(f"sqlite:///{sqlite_path.as_posix()}")

    return create_engine(f"sqlite:///{default_sqlite.as_posix()}")


def main():
    parser = argparse.ArgumentParser(description="Gera uma planilha Excel de cursos e campus por IES.")
    parser.add_argument("sigla", help="Sigla da IES para filtrar (ex: UFFS)")
    args = parser.parse_args()

    sigla = args.sigla.upper()

    sql = """
    SELECT distinct
    i.nome as ies,
    i.sigla,
    cc.ano_censo, 
    case 
        when cc.tp_modalidade_ensino = 2 then c.nome || '-EaD'
        when cc.tp_modalidade_ensino = 1 then c.nome
    end curso,
    m.nome || '/' || m.estado as campus
    from curso_censo cc 
    join curso c on c.codigo = cc.curso 
    join ies i on i.codigo = cc.ies 
    join municipio m on m.codigo = cc.municipio  
    where cc.ano_censo in (2014, 2024) and cc.tp_grau_academico = 2 and i.sigla = :sigla
    """

    engine = carregar_engine()
    print(f"Executando a consulta para a sigla: {sigla}...")
    
    df = pd.read_sql(text(sql), engine, params={"sigla": sigla})
    
    if df.empty:
        print(f"Nenhum dado encontrado para a sigla {sigla}.")
        return

    # Extraindo o nome da IES para o cabeçalho
    ies_nome = df['ies'].iloc[0]

    # Separando os dados por ano e ordenando
    df_2014 = df[df['ano_censo'] == 2014][['curso', 'campus']].sort_values(by=['curso', 'campus']).reset_index(drop=True)
    df_2024 = df[df['ano_censo'] == 2024][['curso', 'campus']].sort_values(by=['curso', 'campus']).reset_index(drop=True)

    # Criação da planilha
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"

    # a) A coluna "ies", deve ser no cabeçalho e não se repetir nas linhas
    ws.merge_cells('A1:D1')
    ws['A1'] = ies_nome
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

    # b) a segunda linha da tabela deverá conter os anos 2014 e 2024 mesclados
    ws.merge_cells('A2:B2')
    ws['A2'] = '2014'
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('C2:D2')
    ws['C2'] = '2024'
    ws['C2'].alignment = Alignment(horizontal="center", vertical="center")

    # c) e d) Linha 3 com os cabeçalhos das colunas
    headers = ['Nome do Curso', 'Campus/Polo', 'Nome do Curso', 'Campus/Polo']
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col_idx, value=header)

    # Preenchendo os dados a partir da linha 4
    max_rows = max(len(df_2014), len(df_2024))
    for i in range(max_rows):
        row_idx = i + 4
        if i < len(df_2014):
            ws.cell(row=row_idx, column=1, value=df_2014.loc[i, 'curso'])
            ws.cell(row=row_idx, column=2, value=df_2014.loc[i, 'campus'])
        if i < len(df_2024):
            ws.cell(row=row_idx, column=3, value=df_2024.loc[i, 'curso'])
            ws.cell(row=row_idx, column=4, value=df_2024.loc[i, 'campus'])

    # Aplicando estilos
    font_default = Font(name='Times New Roman', size=10)
    font_bold = Font(name='Times New Roman', size=10, bold=True)
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_green = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fill_orange = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    fill_grey = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = thin_border
            if cell.row == 1:
                cell.font = font_bold
                cell.fill = fill_yellow
            elif cell.row == 2:
                cell.font = font_bold
                if cell.column <= 2:
                    cell.fill = fill_green
                else:
                    cell.fill = fill_orange
            elif cell.row == 3:
                cell.font = font_bold
                cell.fill = fill_grey
            else:
                cell.font = font_default

    # Ajuste de largura das colunas
    for col in ['A', 'C']:
        ws.column_dimensions[col].width = 40
    for col in ['B', 'D']:
        ws.column_dimensions[col].width = 30

    arquivo_saida = Path(__file__).parent / f"licenciaturas_{sigla}.xlsx"
    wb.save(arquivo_saida)
    
    print(f"Planilha gerada com sucesso: {arquivo_saida}")

if __name__ == "__main__":
    main()
