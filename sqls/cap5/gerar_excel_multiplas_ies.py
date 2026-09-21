import configparser
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from sqlalchemy import create_engine, text

def carregar_engine() -> object:
    root = Path(__file__).resolve().parents[2]
    config_path = root / "config.ini"
    default_sqlite = root / "INEP.db"

    if not config_path.exists():
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    parser = configparser.ConfigParser()
    parser.read(config_path, encoding="utf-8")

    if "BD_INEP" not in parser:
        return create_engine(f"sqlite:///{default_sqlite.as_posix()}")

    bd = parser["BD_INEP"]
    sgdb = bd.get("SGDB", "sqlite").strip().lower()

    if sgdb == "sqlite":
        pasta_local = bd.get("PASTA_LOCAL", "").replace("\\", "/")
        arquivo = bd.get("DW", "INEP.db").strip()
        if pasta_local:
            sqlite_path = Path(pasta_local.strip("/")) / arquivo
        else:
            sqlite_path = default_sqlite

        if not sqlite_path.exists():
            sqlite_path = default_sqlite

        return create_engine(f"sqlite:///{sqlite_path.as_posix()}")

    return create_engine(f"sqlite:///{default_sqlite.as_posix()}")


def main():
    engine = carregar_engine()
    
    print("Buscando as IES de acordo com o filtro...")
    sql_ies = """
    select i.codigo, i.nome, i.sigla 
    from ies i 
    where i.categoria = 1 and i.org_academica = 1 and i.sigla <> '-'
    order by i.sigla
    """
    df_ies = pd.read_sql(text(sql_ies), engine)
    
    if df_ies.empty:
        print("Nenhuma IES encontrada.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    
    font_default = Font(name='Times New Roman', size=10)
    font_bold = Font(name='Times New Roman', size=10, bold=True)
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_green = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fill_orange = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    fill_grey = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    current_row = 1
    
    sql_dados = """
    SELECT distinct
    cc.ano_censo, 
    case 
        when cc.tp_modalidade_ensino = 2 then c.nome || '-EaD'
        when cc.tp_modalidade_ensino = 1 then c.nome
    end curso,
    m.nome || '/' || m.estado as campus
    from curso_censo cc 
    join curso c on c.codigo = cc.curso 
    join municipio m on m.codigo = cc.municipio  
    where cc.ano_censo in (2014, 2024) and cc.tp_grau_academico = 2 and cc.ies = :ies_cod
    """

    for idx, row in df_ies.iterrows():
        ies_cod = row['codigo']
        ies_nome = row['nome']
        ies_sigla = row['sigla']
        
        print(f"Processando IES: {ies_sigla} - {ies_nome}")
        
        df = pd.read_sql(text(sql_dados), engine, params={"ies_cod": int(ies_cod)})
        
        if df.empty:
            continue

        df_2014 = df[df['ano_censo'] == 2014][['curso', 'campus']].sort_values(by=['curso', 'campus']).reset_index(drop=True)
        df_2024 = df[df['ano_censo'] == 2024][['curso', 'campus']].sort_values(by=['curso', 'campus']).reset_index(drop=True)

        start_row = current_row

        # a) A coluna "ies" no cabeçalho (mesclada para a IES atual)
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=4)
        cell_ies = ws.cell(row=start_row, column=1, value=ies_nome)
        cell_ies.alignment = Alignment(horizontal="center", vertical="center")
        cell_ies.font = font_bold
        cell_ies.fill = fill_yellow
        
        for col in range(1, 5):
            ws.cell(row=start_row, column=col).border = thin_border

        # b) a segunda linha da tabela (2014 e 2024)
        ws.merge_cells(start_row=start_row+1, start_column=1, end_row=start_row+1, end_column=2)
        cell_2014 = ws.cell(row=start_row+1, column=1, value='2014')
        cell_2014.alignment = Alignment(horizontal="center", vertical="center")
        cell_2014.font = font_bold
        cell_2014.fill = fill_green
        
        ws.cell(row=start_row+1, column=2).border = thin_border

        ws.merge_cells(start_row=start_row+1, start_column=3, end_row=start_row+1, end_column=4)
        cell_2024 = ws.cell(row=start_row+1, column=3, value='2024')
        cell_2024.alignment = Alignment(horizontal="center", vertical="center")
        cell_2024.font = font_bold
        cell_2024.fill = fill_orange

        for col in range(1, 5):
            ws.cell(row=start_row+1, column=col).border = thin_border

        # c) e d) Linha 3 com os cabeçalhos das colunas
        headers = ['Nome do Curso', 'Campus/Polo', 'Nome do Curso', 'Campus/Polo']
        for col_idx, header in enumerate(headers, start=1):
            cell_header = ws.cell(row=start_row+2, column=col_idx, value=header)
            cell_header.font = font_bold
            cell_header.fill = fill_grey
            cell_header.border = thin_border

        # Preenchendo os dados da IES
        max_rows = max(len(df_2014), len(df_2024))
        for i in range(max_rows):
            row_idx = start_row + 3 + i
            
            # Dados de 2014
            if i < len(df_2014):
                c1 = ws.cell(row=row_idx, column=1, value=df_2014.loc[i, 'curso'])
                c2 = ws.cell(row=row_idx, column=2, value=df_2014.loc[i, 'campus'])
            else:
                c1 = ws.cell(row=row_idx, column=1)
                c2 = ws.cell(row=row_idx, column=2)
            c1.font = font_default; c1.border = thin_border
            c2.font = font_default; c2.border = thin_border

            # Dados de 2024
            if i < len(df_2024):
                c3 = ws.cell(row=row_idx, column=3, value=df_2024.loc[i, 'curso'])
                c4 = ws.cell(row=row_idx, column=4, value=df_2024.loc[i, 'campus'])
            else:
                c3 = ws.cell(row=row_idx, column=3)
                c4 = ws.cell(row=row_idx, column=4)
            c3.font = font_default; c3.border = thin_border
            c4.font = font_default; c4.border = thin_border
                
        # Adiciona 1 linha vazia mesclada entre as IES
        empty_row = start_row + 3 + max_rows
        ws.merge_cells(start_row=empty_row, start_column=1, end_row=empty_row, end_column=4)
        current_row = empty_row + 1
    # Ajuste de largura das colunas
    for col in ['A', 'C']:
        ws.column_dimensions[col].width = 40
    for col in ['B', 'D']:
        ws.column_dimensions[col].width = 30

    arquivo_saida = Path(__file__).parent / "licenciaturas_todas_ies.xlsx"
    wb.save(arquivo_saida)
    
    print(f"Planilha geral gerada com sucesso: {arquivo_saida}")

if __name__ == "__main__":
    main()
